# -*- coding: utf-8 -*-
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

import os
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from functools import wraps
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pandas as pd
import psycopg2
from psycopg2.extras import RealDictCursor
from flask import (
    Flask, flash, jsonify, redirect, render_template,
    request, send_file, url_for, session
)
from dotenv import load_dotenv
from werkzeug.utils import secure_filename

# Charger les variables d'environnement
import pathlib
env_file = pathlib.Path(__file__).parent / '.env'
if env_file.exists():
    with open(env_file, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith('#'):
                key, value = line.split('=', 1)
                os.environ[key.strip()] = value.strip()

# Fallback pour charger_dotenv
load_dotenv()

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "annuaire_quartier_2025_secret")
app.permanent_session_lifetime = timedelta(minutes=45)

import bcrypt
#for user, pw in [('admin','1234'),('agent','0000'),('responsable','1111'),('hote','2222')]:
    #h = bcrypt.hashpw(pw.encode(), bcrypt.gensalt()).decode()
    #print(f"UPDATE utilisateurs SET password='{h}' WHERE username='{user}';")

import random
import string
from flask_mail import Mail, Message as MailMessage

# Configuration Flask-Mail (Gmail SMTP)
app.config['MAIL_SERVER']   = 'smtp.gmail.com'
app.config['MAIL_PORT']     = 587
app.config['MAIL_USE_TLS']  = True
app.config['MAIL_USERNAME'] = 'mfad09012002@gmail.com'
app.config['MAIL_PASSWORD'] = os.environ.get('MAIL_PASSWORD', '')
app.config['MAIL_DEFAULT_SENDER'] = ('Annuaire Quartier', 'mfad09012002@gmail.com')

mail = Mail(app)

reset_codes = {}

# Connexion DB
def get_db():
    """Établit une connexion à PostgreSQL avec gestion d'encodage UTF-8."""
    try:
        from urllib.parse import quote

        database_url = os.environ.get("DATABASE_URL")
        if database_url:
            return psycopg2.connect(database_url, client_encoding="UTF8")
        
        # Récupère les variables depuis l'environnement
        host = os.environ.get("DB_HOST", "localhost")
        port = os.environ.get("DB_PORT", "5432")
        database = os.environ.get("DB_NAME", "bd_quartier")
        user = os.environ.get("DB_USER", "postgres")
        password = os.environ.get("DB_PASSWORD", "")
        
        # Encode les paramètres correctement
        password_encoded = quote(str(password), safe='')
        
        # Utilise une URL de connexion pour éviter les problèmes d'encodage
        dsn = f"postgresql://{user}:{password_encoded}@{host}:{port}/{database}?client_encoding=UTF8"
        
        conn = psycopg2.connect(dsn)
        return conn
    except Exception as e:
        print(f"Erreur de connexion DB: {e}")
        raise

# Enregistrement des actions dans l'historique
def log_action(action, details=""):
    """Enregistre une action dans la table historique_actions."""
    try:
        conn = get_db()
        cur  = conn.cursor()
        cur.execute("""
            INSERT INTO historique_actions (utilisateur, role, action, details, date_action)
            VALUES (%s, %s, %s, %s, NOW())
        """, (
            session.get("user", "système"),
            session.get("role", "—"),
            action,
            details,
        ))
        conn.commit()
        cur.close(); conn.close()
    except Exception:
        pass  # Ne pas bloquer l'appel principal si le log échoue
# ─────────────────────────────────────────────────────────────────
# Décorateurs d'accès
# ─────────────────────────────────────────────────────────────────
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user" not in session:
            flash("Veuillez vous connecter.", "warning")
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated

def roles_required(*roles):
    def decorator(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            if "user" not in session:
                return redirect(url_for("login"))
            if session.get("role") not in roles:
                flash("Accès non autorisé pour votre profil.", "danger")
                return redirect(url_for("accueil"))
            return f(*args, **kwargs)
        return decorated
    return decorator




def valider_menage(adresse, revenu, latitude, longitude, score):
    erreurs = []

    # 1. Adresse obligatoire
    if not adresse or adresse.strip() == "":
        erreurs.append("L'adresse du ménage est obligatoire.")

    # 2. Revenu positif ou nul
    if revenu not in (None, ""):
        try:
            revenu_val = int(revenu)
            if revenu_val < 0:
                erreurs.append("Le revenu ne peut pas être négatif.")
        except ValueError:
            erreurs.append("Le revenu doit être un nombre.")

    # 3. Score entre 0 et 100
    if score not in (None, ""):
        try:
            score_val = float(score)
            if score_val < 0 or score_val > 100:
                erreurs.append("Le score de vulnérabilité doit être compris entre 0 et 100.")
        except ValueError:
            erreurs.append("Le score de vulnérabilité doit être un nombre.")

    # 4. Latitude et longitude doivent aller ensemble
    if bool(latitude) != bool(longitude):
        erreurs.append("La latitude et la longitude doivent être renseignées ensemble.")

    # 5. Vérifier les coordonnées
    if latitude and longitude:
        try:
            lat = float(latitude)
            lon = float(longitude)

            if lat < -90 or lat > 90:
                erreurs.append("La latitude doit être comprise entre -90 et 90.")

            if lon < -180 or lon > 180:
                erreurs.append("La longitude doit être comprise entre -180 et 180.")

            # Zone approximative d'Abidjan
            if not (5.20 <= lat <= 5.50 and -4.20 <= lon <= -3.80):
                erreurs.append("Les coordonnées semblent hors de la zone d'Abidjan.")

        except ValueError:
            erreurs.append("Les coordonnées doivent être numériques.")

    return erreurs


def normalize_key(value):
    """Normalise un libelle pour mapper proprement les colonnes importees."""
    value = str(value or "").strip().lower()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "_", value).strip("_")


def safe_int(value):
    if value in (None, "") or pd.isna(value):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def clean_text(value):
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def parse_import_date(value):
    if value in (None, "") or pd.isna(value):
        return None
    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        return None
    return parsed.date()


def normalize_sexe(value):
    v = normalize_key(value)
    if v in ("m", "masculin", "homme", "garcon"):
        return "Masculin"
    if v in ("f", "feminin", "femme", "fille"):
        return "Féminin"
    return clean_text(value)


def normalize_statut(value):
    v = normalize_key(value)
    if not v:
        return "actif"
    if v in ("actif", "active", "present", "vivant"):
        return "actif"
    if v in ("decede", "deces", "mort"):
        return "decede"
    if v in ("parti", "depart", "sorti"):
        return "parti"
    return clean_text(value)


def validate_password_strength(password):
    errors = []
    if len(password or "") < 8:
        errors.append("Le mot de passe doit contenir au moins 8 caractères.")
    if password and not re.search(r"[A-Za-z]", password):
        errors.append("Le mot de passe doit contenir au moins une lettre.")
    if password and not re.search(r"\d", password):
        errors.append("Le mot de passe doit contenir au moins un chiffre.")
    return errors


def collect_report_stats(cur):
    stats = {}
    for key, sql in [
        ("pop_active", "SELECT COUNT(*) c FROM Habitant WHERE statut='actif'"),
        ("pop_totale", "SELECT COUNT(*) c FROM Habitant"),
        ("nb_decedes", "SELECT COUNT(*) c FROM Habitant WHERE statut='decede'"),
        ("nb_partis", "SELECT COUNT(*) c FROM Habitant WHERE statut='parti'"),
        ("nb_menages", "SELECT COUNT(*) c FROM Menages"),
        ("total_naissances", "SELECT COUNT(*) c FROM Evenement_vital WHERE type_evenement='naissance'"),
        ("total_deces", "SELECT COUNT(*) c FROM Evenement_vital WHERE type_evenement='deces'"),
        ("total_arrivees", "SELECT COUNT(*) c FROM Mouvement_residuel WHERE type_mouvement='arrivee'"),
        ("total_departs", "SELECT COUNT(*) c FROM Mouvement_residuel WHERE type_mouvement='depart'"),
    ]:
        cur.execute(sql)
        stats[key] = cur.fetchone()["c"]

    cur.execute("""
        SELECT ROUND(AVG(nb),2) as moy FROM
        (SELECT COUNT(*) nb FROM Habitant
         WHERE statut = 'actif' GROUP BY id_menage) t
    """)
    row = cur.fetchone()
    stats["taille_moy_menage"] = row["moy"] if row and row["moy"] else 0
    return stats


def build_quality_snapshot(cur):
    metrics = {}
    checks = [
        ("habitants_sans_telephone", "SELECT COUNT(*) c FROM Habitant WHERE COALESCE(TRIM(telephone),'') = ''"),
        ("habitants_sans_naissance", "SELECT COUNT(*) c FROM Habitant WHERE date_de_naissance IS NULL"),
        ("habitants_sans_menage", "SELECT COUNT(*) c FROM Habitant WHERE id_menage IS NULL"),
        ("telephones_invalides", "SELECT COUNT(*) c FROM Habitant WHERE COALESCE(TRIM(telephone),'') <> '' AND telephone !~ '^[+0-9 ()-]{8,20}$'"),
        ("ages_anormaux", "SELECT COUNT(*) c FROM Habitant WHERE date_de_naissance > CURRENT_DATE OR EXTRACT(YEAR FROM AGE(date_de_naissance)) > 120"),
        ("menages_sans_chef", "SELECT COUNT(*) c FROM Menages m WHERE m.id_chef_menage IS NULL OR NOT EXISTS (SELECT 1 FROM Habitant h WHERE h.id_habitant=m.id_chef_menage AND h.statut='actif')"),
        ("menages_sans_coordonnees", "SELECT COUNT(*) c FROM Menages WHERE latitude IS NULL OR longitude IS NULL"),
        ("menages_vulnerables", "SELECT COUNT(*) c FROM Menages WHERE COALESCE(score_vulnerabilite,0) >= 70"),
        ("deces_sans_statut", "SELECT COUNT(*) c FROM Evenement_vital e JOIN Habitant h ON h.id_habitant=e.id_habitant WHERE e.type_evenement='deces' AND h.statut <> 'decede'"),
    ]
    for key, sql in checks:
        cur.execute(sql)
        metrics[key] = cur.fetchone()["c"]

    cur.execute("""
        SELECT UPPER(TRIM(nom)) nom, UPPER(TRIM(prenom)) prenom, date_de_naissance, COUNT(*) total
        FROM Habitant
        GROUP BY UPPER(TRIM(nom)), UPPER(TRIM(prenom)), date_de_naissance
        HAVING COUNT(*) > 1
        ORDER BY total DESC, nom, prenom
        LIMIT 10
    """)
    duplicates = cur.fetchall()
    metrics["doublons_potentiels"] = sum(row["total"] for row in duplicates) if duplicates else 0

    cur.execute("""
        SELECT h.id_habitant, h.nom, h.prenom, h.telephone, h.date_de_naissance, h.statut, m.adresse
        FROM Habitant h
        LEFT JOIN Menages m ON m.id_menage = h.id_menage
        WHERE COALESCE(TRIM(h.telephone),'') = ''
           OR h.date_de_naissance IS NULL
           OR h.id_menage IS NULL
           OR h.date_de_naissance > CURRENT_DATE
           OR EXTRACT(YEAR FROM AGE(h.date_de_naissance)) > 120
        ORDER BY h.nom, h.prenom
        LIMIT 25
    """)
    habitants_a_corriger = cur.fetchall()

    cur.execute("""
        SELECT m.id_menage, m.adresse, m.secteur, m.score_vulnerabilite,
               COUNT(h.id_habitant) nb_membres,
               CASE WHEN m.id_chef_menage IS NULL THEN TRUE ELSE FALSE END sans_chef
        FROM Menages m
        LEFT JOIN Habitant h ON h.id_menage = m.id_menage AND h.statut='actif'
        WHERE m.id_chef_menage IS NULL
           OR m.latitude IS NULL
           OR m.longitude IS NULL
           OR COALESCE(m.score_vulnerabilite,0) >= 70
        GROUP BY m.id_menage
        ORDER BY COALESCE(m.score_vulnerabilite,0) DESC, m.adresse
        LIMIT 25
    """)
    menages_a_corriger = cur.fetchall()

    cur.execute("SELECT COUNT(*) c FROM Habitant")
    total_habitants = cur.fetchone()["c"] or 0
    cur.execute("SELECT COUNT(*) c FROM Menages")
    total_menages = cur.fetchone()["c"] or 0
    issue_total = sum(v for v in metrics.values() if isinstance(v, int))
    denominator = max(total_habitants + total_menages, 1)
    score = max(0, min(100, round(100 - (issue_total / denominator * 18))))

    return {
        "metrics": metrics,
        "score": score,
        "issue_total": issue_total,
        "duplicates": duplicates,
        "habitants_a_corriger": habitants_a_corriger,
        "menages_a_corriger": menages_a_corriger,
    }

# ─────────────────────────────────────────────────────────────────
# Auth
# ─────────────────────────────────────────────────────────────────

# Dans app.py — AVANT la route /login

@app.route("/")
def home():
    if "user" in session:
        return redirect(url_for('accueil'))
    return redirect(url_for("login"))

@app.route("/accueil")
@login_required
def accueil():
    return render_template("accueil.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if "user" in session:
        return redirect(url_for("accueil"))

    # Charger dynamiquement la liste des profils
    conn = get_db(); cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("SELECT username, label FROM utilisateurs ORDER BY label")
    profils = cur.fetchall()
    cur.close(); conn.close()

    if request.method == "POST":
        username = request.form.get("role", "").strip()
        password = request.form.get("password", "").encode()
        conn = get_db(); cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute("SELECT * FROM utilisateurs WHERE username = %s", (username,))
        user = cur.fetchone()
        cur.close(); conn.close()
        if user and bcrypt.checkpw(password, user["password"].encode()):
            session.permanent = True
            session["user"]  = user["username"]
            session["role"]  = user["role"]
            session["label"] = user["label"]
            log_action("Connexion", f"Profil : {user['username']}")
            flash(f"Bienvenue, {user['label']} !", "success")
            return redirect(url_for("accueil"))
        else:
            flash("Identifiants incorrects. Réessayez.", "danger")
    return render_template("login.html", profils=profils)
@app.route("/logout")
def logout():
    log_action("Déconnexion", f"Profil : {session.get('user','?')}")
    session.clear()
    flash("Vous avez été déconnecté.", "info")
    return redirect(url_for("login"))
# ─────────────────────────────────────────────────────────────────
# Mot de passe oublié
# ─────────────────────────────────────────────────────────────────
@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    conn = get_db(); cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("SELECT username, label FROM utilisateurs ORDER BY label")
    utilisateurs = cur.fetchall()
    cur.close(); conn.close()

    if request.method == "POST":
        username    = request.form.get("role", "").strip()
        email_input = request.form.get("email", "").strip()
        conn = get_db(); cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute("SELECT * FROM utilisateurs WHERE username=%s", (username,))
        user = cur.fetchone()
        cur.close(); conn.close()
        if not user:
            flash("Profil inconnu.", "danger")
            return render_template("forgot_password.html", utilisateurs=utilisateurs)
        if email_input != user["email"]:
            flash("Email incorrect pour ce profil.", "danger")
            return render_template("forgot_password.html", utilisateurs=utilisateurs)
        code = ''.join(random.choices(string.digits, k=6))
        reset_codes[username] = code
        try:
            msg = MailMessage(subject="Code de réinitialisation — Annuaire Quartier",
                              recipients=[email_input],
                              html=f"<h2>Votre code : {code}</h2>")
            mail.send(msg)
            flash("Code envoyé à votre adresse email.", "success")
        except Exception as e:
            flash(f"Erreur d'envoi email : {e}", "danger")
            return render_template("forgot_password.html", utilisateurs=utilisateurs)
        return redirect(url_for("reset_password", role=username))
    return render_template("forgot_password.html", utilisateurs=utilisateurs)

@app.route("/reset-password/<role>", methods=["GET", "POST"])
def reset_password(role):
    conn = get_db(); cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("SELECT * FROM utilisateurs WHERE username=%s", (role,))
    user = cur.fetchone(); cur.close(); conn.close()
    if not user:
        flash("Profil invalide.", "danger")
        return redirect(url_for("login"))
    if request.method == "POST":
        code_saisi       = request.form.get("code", "").replace(" ", "").strip()
        new_password     = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")
        if reset_codes.get(role) != code_saisi:
            flash("Code incorrect.", "danger")
            return render_template("reset_password.html", role=role)
        if new_password != confirm_password:
            flash("Les mots de passe ne correspondent pas.", "danger")
            return render_template("reset_password.html", role=role)
        if len(new_password) < 6:
            flash("Mot de passe trop court (minimum 6 caractères).", "danger")
            return render_template("reset_password.html", role=role)
        hashed = bcrypt.hashpw(new_password.encode(), bcrypt.gensalt()).decode()
        conn = get_db(); cur = conn.cursor()
        cur.execute("UPDATE utilisateurs SET password=%s WHERE username=%s", (hashed, role))
        conn.commit(); cur.close(); conn.close()
        reset_codes.pop(role, None)
        flash("Mot de passe modifié avec succès.", "success")
        return redirect(url_for("login"))
    return render_template("reset_password.html", role=role)
# ─────────────────────────────────────────────────────────────────
# Tableau de bord
# ─────────────────────────────────────────────────────────────────
@app.route("/dashboard")
@login_required
def dashboard():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    stats = {}
        # KPI globaux du rapport
    cur.execute("SELECT COUNT(*) as total FROM Habitant WHERE statut='decede'")
    stats["nb_decedes"] = cur.fetchone()["total"]

    cur.execute("SELECT COUNT(*) as total FROM Habitant WHERE statut='parti'")
    stats["nb_partis"] = cur.fetchone()["total"]

    cur.execute("SELECT COUNT(*) as total FROM Evenement_vital WHERE type_evenement='naissance'")
    stats["total_naissances"] = cur.fetchone()["total"]

    cur.execute("SELECT COUNT(*) as total FROM Evenement_vital WHERE type_evenement='deces'")
    stats["total_deces"] = cur.fetchone()["total"]

    cur.execute("SELECT COUNT(*) as total FROM Mouvement_residuel WHERE type_mouvement='arrivee'")
    stats["total_arrivees"] = cur.fetchone()["total"]

    cur.execute("SELECT COUNT(*) as total FROM Mouvement_residuel WHERE type_mouvement='depart'")
    stats["total_departs"] = cur.fetchone()["total"]

    cur.execute("""
        SELECT ROUND(AVG(nb),2) as moy FROM
        (
            SELECT COUNT(*) nb 
            FROM Habitant
            WHERE statut='actif'
            GROUP BY id_menage
        ) AS t
    """)
    row = cur.fetchone()
    stats["taille_moy_menage"] = row["moy"] if row and row["moy"] else 0
    cur.execute("SELECT COUNT(*) as total FROM Habitant WHERE statut = 'actif'")
    stats["population"] = cur.fetchone()["total"]

    cur.execute("SELECT COUNT(*) as total FROM Habitant")
    stats["population_totale"] = cur.fetchone()["total"]

    cur.execute("SELECT COUNT(*) as total FROM Menages")
    stats["menages"] = cur.fetchone()["total"]

    cur.execute("""
        SELECT COUNT(*) as total FROM Evenement_vital
        WHERE type_evenement = 'naissance'
        AND DATE_TRUNC('month', date_evenement) = DATE_TRUNC('month', CURRENT_DATE)
    """)
    stats["naissances_mois"] = cur.fetchone()["total"]

    cur.execute("""
        SELECT COUNT(*) as total FROM Evenement_vital
        WHERE type_evenement = 'deces'
        AND DATE_TRUNC('month', date_evenement) = DATE_TRUNC('month', CURRENT_DATE)
    """)
    stats["deces_mois"] = cur.fetchone()["total"]

    cur.execute("""
        SELECT COUNT(*) as total FROM Mouvement_residuel
        WHERE type_mouvement = 'arrivee'
        AND DATE_TRUNC('month', date_mouvement) = DATE_TRUNC('month', CURRENT_DATE)
    """)
    stats["arrivees_mois"] = cur.fetchone()["total"]

    cur.execute("""
        SELECT COUNT(*) as total FROM Mouvement_residuel
        WHERE type_mouvement = 'depart'
        AND DATE_TRUNC('month', date_mouvement) = DATE_TRUNC('month', CURRENT_DATE)
    """)
    stats["departs_mois"] = cur.fetchone()["total"]

    cur.execute("""
        SELECT ROUND(AVG(
                EXTRACT(YEAR FROM AGE(CURRENT_DATE, date_de_naissance))
            ), 2) AS total
        FROM Habitant
        WHERE statut = 'actif';
        """)
    stats["age_moyenne"] = cur.fetchone()["total"]

    # Répartition par sexe
    cur.execute("""
        SELECT COALESCE(NULLIF(sexe,''), 'Non renseigné') as sexe, COUNT(*) as count
        FROM Habitant WHERE statut = 'actif'
        GROUP BY sexe ORDER BY count DESC
    """)
    sexe_data = cur.fetchall()

    # Répartition par tranche d'âge
    cur.execute("""
        SELECT
            CASE
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 0  AND 9  THEN '0–9'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 10 AND 19 THEN '10–19'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 20 AND 29 THEN '20–29'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 30 AND 39 THEN '30–39'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 40 AND 49 THEN '40–49'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 50 AND 59 THEN '50–59'
                ELSE '60+'
            END as tranche,
            COUNT(*) as count
        FROM Habitant
        WHERE date_de_naissance IS NOT NULL AND (statut = 'actif')
        GROUP BY tranche ORDER BY tranche DESC
    """)
    age_data = cur.fetchall()

    # Évolution journaliere deces/naissances
    cur.execute("""
        SELECT
            TO_CHAR(DATE_TRUNC('day', date_evenement), 'DD Mon YYYY') AS jour,
            DATE_TRUNC('day', date_evenement) AS jour_order,
            SUM(
                CASE 
                    WHEN type_evenement = 'naissance' THEN 1 
                    ELSE 0 
                END
            ) AS naissances,
            SUM(
                CASE 
                    WHEN type_evenement = 'deces' THEN 1 
                    ELSE 0 
                END
            ) AS deces
        FROM Evenement_vital
        GROUP BY DATE_TRUNC('day', date_evenement)
        ORDER BY jour_order;
    """)
    evolution_data = cur.fetchall()

    # Récents événements
    cur.execute("""
        SELECT e.id_evenement, e.type_evenement, e.date_evenement, e.description,
               h.nom, h.prenom
        FROM Evenement_vital e
        JOIN Habitant h ON e.id_habitant = h.id_habitant
        ORDER BY e.date_evenement DESC LIMIT 6
    """)
    recent_events = cur.fetchall()

    # Récents habitants ajoutés
    cur.execute("""
        SELECT h.*, m.adresse FROM Habitant h
        LEFT JOIN Menages m ON h.id_menage = m.id_menage
        ORDER BY h.id_habitant DESC LIMIT 5
    """)
    recent_habitants = cur.fetchall()
        # Top professions
    cur.execute("""
        SELECT COALESCE(NULLIF(profession,''),'Non renseigné') as profession,
            COUNT(*) as cnt
        FROM Habitant
        GROUP BY profession
        ORDER BY cnt DESC
        LIMIT 8
    """)
    professions = cur.fetchall()

    # Niveaux d'étude
    cur.execute("""
        SELECT COALESCE(NULLIF(niveau_etude,''),'Non renseigné') as niveau_etude,
            COUNT(*) as cnt
        FROM Habitant
        GROUP BY niveau_etude
        ORDER BY cnt DESC
    """)
    niveaux = cur.fetchall()

    # Évolution 12 mois
    cur.execute("""
        SELECT
            TO_CHAR(DATE_TRUNC('month', date_evenement), 'Mon YYYY') as mois,
            DATE_TRUNC('month', date_evenement) as mois_order,
            SUM(CASE WHEN type_evenement='naissance' THEN 1 ELSE 0 END) as naissances,
            SUM(CASE WHEN type_evenement='deces' THEN 1 ELSE 0 END) as deces
        FROM Evenement_vital
        WHERE date_evenement >= CURRENT_DATE - INTERVAL '12 months'
        GROUP BY DATE_TRUNC('month', date_evenement)
        ORDER BY mois_order
    """)
    evolution_rapport = cur.fetchall()
    cur.close(); conn.close()
    return render_template(
        "dashboard.html",
        stats=stats,
        sexe_data=sexe_data,
        age_data=age_data,
        evolution_data=evolution_data,
        professions=professions,
        niveaux=niveaux,
        evolution_rapport=evolution_rapport,
        recent_events=recent_events,
        recent_habitants=recent_habitants,
        today=date.today().strftime('%d/%m/%Y'),
    )

# ─────────────────────────────────────────────────────────────────
# Habitants
# ─────────────────────────────────────────────────────────────────
@app.route("/habitants")
@login_required
def habitants_list():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    q = request.args.get("q", "").strip()
    sexe = request.args.get("sexe", "")
    statut = request.args.get("statut", "")
    profession = request.args.get("profession", "").strip()
    niveau = request.args.get("niveau", "").strip()
    secteur = request.args.get("secteur", "").strip()
    age_min = request.args.get("age_min", "").strip()
    age_max = request.args.get("age_max", "").strip()
    sans_menage = request.args.get("sans_menage", "")

    sql = """
        SELECT h.*, m.adresse, m.secteur,
               CASE
                 WHEN h.date_de_naissance IS NULL THEN NULL
                 ELSE EXTRACT(YEAR FROM AGE(h.date_de_naissance))::INT
               END AS age
        FROM Habitant h
        LEFT JOIN Menages m ON h.id_menage = m.id_menage
        WHERE 1=1
    """
    params = []
    if q:
        sql += " AND (h.nom ILIKE %s OR h.prenom ILIKE %s OR h.telephone ILIKE %s OR m.adresse ILIKE %s)"
        params += [f"%{q}%"] * 4
    if sexe:
        sql += " AND h.sexe = %s"; params.append(sexe)
    if statut:
        sql += " AND h.statut = %s"; params.append(statut)
    if profession:
        sql += " AND h.profession ILIKE %s"; params.append(f"%{profession}%")
    if niveau:
        sql += " AND h.niveau_etude = %s"; params.append(niveau)
    if secteur:
        sql += " AND m.secteur = %s"; params.append(secteur)
    if safe_int(age_min) is not None:
        sql += " AND h.date_de_naissance IS NOT NULL AND EXTRACT(YEAR FROM AGE(h.date_de_naissance)) >= %s"
        params.append(safe_int(age_min))
    if safe_int(age_max) is not None:
        sql += " AND h.date_de_naissance IS NOT NULL AND EXTRACT(YEAR FROM AGE(h.date_de_naissance)) <= %s"
        params.append(safe_int(age_max))
    if sans_menage:
        sql += " AND h.id_menage IS NULL"
    sql += " ORDER BY h.nom, h.prenom"

    cur.execute(sql, params)
    habitants = cur.fetchall()
    cur.execute("SELECT DISTINCT profession FROM Habitant WHERE COALESCE(profession,'') <> '' ORDER BY profession")
    professions = [r["profession"] for r in cur.fetchall()]
    cur.execute("SELECT DISTINCT niveau_etude FROM Habitant WHERE COALESCE(niveau_etude,'') <> '' ORDER BY niveau_etude")
    niveaux = [r["niveau_etude"] for r in cur.fetchall()]
    cur.execute("SELECT DISTINCT secteur FROM Menages WHERE COALESCE(secteur,'') <> '' ORDER BY secteur")
    secteurs = [r["secteur"] for r in cur.fetchall()]
    cur.close(); conn.close()
    return render_template("habitants_list.html",
                           habitants=habitants, q=q, sexe=sexe, statut=statut,
                           profession=profession, niveau=niveau, secteur=secteur,
                           age_min=age_min, age_max=age_max, sans_menage=sans_menage,
                           professions=professions, niveaux=niveaux, secteurs=secteurs)

@app.route("/habitants/add", methods=["GET", "POST"])
@roles_required("admin", "agent")
def habitant_add():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("SELECT id_menage, adresse FROM Menages ORDER BY adresse")
    menages = cur.fetchall()
    id_menage_preselect = request.args.get("id_menage")
    lien_avec_chef = request.form.get("lien_avec_chef")
    id_menage = request.form.get("id_menage")
    if request.method == "POST":
        d = {
            "nom":              request.form.get("nom", "").strip().upper(),
            "prenom":           request.form.get("prenom", "").strip(),
            "sexe":             request.form.get("sexe", ""),
            "date_de_naissance": request.form.get("date_de_naissance") or None,
            "profession":       request.form.get("profession", "").strip(),
            "niveau_etude":     request.form.get("niveau_etude", "").strip(),
            "telephone":        request.form.get("telephone", "").strip(),
            "statut":           request.form.get("statut", "actif"),
            "id_menage":        request.form.get("id_menage") or None,
            "lien_avec_chef":    request.form.get("lien_avec_chef")
        }
        # Validation serveur
        errors = []
        if not d["nom"]:
            errors.append("Le nom est obligatoire.")
        if not d["prenom"]:
            errors.append("Le prénom est obligatoire.")
        if d["sexe"] and d["sexe"] not in ("Masculin", "Féminin"):
            errors.append("Sexe invalide.")
        if d["statut"] not in ("actif", "parti", "decede"):
            errors.append("Statut invalide.")
        if d["date_de_naissance"]:
            try:
                ddn = datetime.strptime(d["date_de_naissance"], "%Y-%m-%d").date()
                if ddn > date.today():
                    errors.append("La date de naissance ne peut pas être dans le futur.")
            except ValueError:
                errors.append("Date de naissance invalide.")
        if errors:
            for e in errors:
                flash(e, "danger")
            cur.close(); conn.close()
            return render_template(
                            "habitants_add.html",
                            menages=menages,
                            today=date.today().strftime("%Y-%m-%d"),
                            id_menage_preselect=id_menage_preselect
                        )
        
        if lien_avec_chef == "Chef de ménage":
            cur.execute("""
                SELECT id_habitant
                FROM Habitant
                WHERE id_menage = %s
                AND lien_avec_chef = 'Chef de ménage'
                AND statut = 'actif'
            """, (id_menage,))

            chef_existant = cur.fetchone()

            if chef_existant:
                flash("Ce ménage possède déjà un chef de ménage actif.", "danger")
                return render_template(
                    "habitants_add.html",
                    menages=menages,
                    today=date.today().strftime("%Y-%m-%d")
                )
        cur.execute("""
            INSERT INTO Habitant
              (nom, prenom, sexe, date_de_naissance, profession, niveau_etude,
               telephone, statut, id_menage,lien_avec_chef)
            VALUES
              (%(nom)s, %(prenom)s, %(sexe)s, %(date_de_naissance)s, %(profession)s,
               %(niveau_etude)s, %(telephone)s, %(statut)s, %(id_menage)s,%(lien_avec_chef)s)
            RETURNING id_habitant
        """, d)
        new_id = cur.fetchone()["id_habitant"]   # avant commit()
        if d["lien_avec_chef"] == "Chef de ménage" and d["id_menage"]:
            cur.execute("""
                UPDATE Menages
                SET id_chef_menage = %s
                WHERE id_menage = %s
            """, (new_id, d["id_menage"]))
        conn.commit()
        log_action("Ajout habitant", f"{d['prenom']} {d['nom']} (ID {new_id})")
        flash(f"Habitant {d['prenom']} {d['nom']} ajouté avec succès.", "success")
        cur.close(); conn.close()
        return redirect(url_for("habitant_detail", id=new_id))

    cur.close(); conn.close()
    return render_template(
                                "habitants_add.html",
                                menages=menages,
                                today=date.today().strftime("%Y-%m-%d"),
                                id_menage_preselect=id_menage_preselect
                            )

@app.route("/habitants/<int:id>")
@login_required
def habitant_detail(id):
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("""
        SELECT h.*, m.adresse, m.revenu FROM Habitant h
        LEFT JOIN Menages m ON h.id_menage = m.id_menage
        WHERE h.id_habitant = %s
    """, (id,))
    habitant = cur.fetchone()
    if not habitant:
        flash("Habitant introuvable.", "danger")
        return redirect(url_for("habitants_list"))

    age = None
    if habitant["date_de_naissance"]:
        today = date.today()
        b = habitant["date_de_naissance"]
        age = today.year - b.year - ((today.month, today.day) < (b.month, b.day))

    cur.execute("""
        SELECT * FROM Evenement_vital WHERE id_habitant=%s ORDER BY date_evenement DESC
    """, (id,))
    evenements = cur.fetchall()

    cur.execute("""
        SELECT * FROM Mouvement_residuel WHERE id_habitant=%s ORDER BY date_mouvement DESC
    """, (id,))
    mouvements = cur.fetchall()

    cur.close(); conn.close()
    return render_template("habitants_detail.html",
                           habitant=habitant, age=age,
                           evenements=evenements, mouvements=mouvements)

@app.route("/habitants/<int:id>/edit", methods=["GET", "POST"])
@roles_required("admin", "agent")
def habitant_edit(id):
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("SELECT * FROM Habitant WHERE id_habitant=%s", (id,))
    habitant = cur.fetchone()
    if not habitant:
        flash("Habitant introuvable.", "danger")
        return redirect(url_for("habitants_list"))
    cur.execute("SELECT id_menage, adresse FROM Menages ORDER BY adresse")
    menages = cur.fetchall()

    if request.method == "POST":
        d = {
            "id":               id,
            "nom":              request.form.get("nom", "").strip().upper(),
            "prenom":           request.form.get("prenom", "").strip(),
            "sexe":             request.form.get("sexe", ""),
            "date_de_naissance": request.form.get("date_de_naissance") or None,
            "profession":       request.form.get("profession", "").strip(),
            "niveau_etude":     request.form.get("niveau_etude", "").strip(),
            "telephone":        request.form.get("telephone", "").strip(),
            "statut":           request.form.get("statut", "actif"),
            "id_menage":        request.form.get("id_menage") or None,
            "lien_avec_chef":    request.form.get("lien_avec_chef"),
        }
        # Validation serveur

        errors = []
        if d["lien_avec_chef"] == "Chef de ménage":
            if not d["id_menage"]:
                errors.append("Un chef de ménage doit être rattaché à un ménage.")
        if not d["nom"]:
            errors.append("Le nom est obligatoire.")
        if not d["prenom"]:
            errors.append("Le prénom est obligatoire.")
        if d["sexe"] and d["sexe"] not in ("Masculin", "Féminin"):
            errors.append("Sexe invalide.")
        if d["statut"] not in ("actif", "parti", "decede"):
            errors.append("Statut invalide.")
        if d["date_de_naissance"]:
            try:
                ddn = datetime.strptime(d["date_de_naissance"], "%Y-%m-%d").date()
                if ddn > date.today():
                    errors.append("La date de naissance ne peut pas être dans le futur.")
            except ValueError:
                errors.append("Date de naissance invalide.")
        if errors:
            for e in errors:
                flash(e, "danger")
            cur.close(); conn.close()
            return render_template("habitants_edit.html", habitant=habitant, menages=menages)
        cur.execute("""
            UPDATE Habitant
            SET nom=%(nom)s,
                prenom=%(prenom)s,
                sexe=%(sexe)s,
                date_de_naissance=%(date_de_naissance)s,
                profession=%(profession)s,
                niveau_etude=%(niveau_etude)s,
                telephone=%(telephone)s,
                statut=%(statut)s,
                id_menage=%(id_menage)s,
                lien_avec_chef=%(lien_avec_chef)s
            WHERE id_habitant=%(id)s
        """, d)

        if d["lien_avec_chef"] == "Chef de ménage" and d["id_menage"]:
            cur.execute("""
                UPDATE Menages
                SET id_chef_menage = %s
                WHERE id_menage = %s
            """, (id, d["id_menage"]))
        else:
            cur.execute("""
                UPDATE Menages
                SET id_chef_menage = NULL
                WHERE id_chef_menage = %s
            """, (id,))
        conn.commit()
        # Ajouter l'action a l'historique

        log_action("Modification habitant", f"ID {id} — {d['prenom']} {d['nom']}")
        flash("Habitant mis à jour avec succès.", "success")
        cur.close(); conn.close()
        return redirect(url_for("habitant_detail", id=id))

    cur.close(); conn.close()
    return render_template("habitants_edit.html", habitant=habitant, menages=menages)

@app.route("/habitants/<int:id>/delete", methods=["POST"])
@roles_required("admin")
def habitant_delete(id):
    conn = get_db()
    cur = conn.cursor()
    try:
        cur.execute("DELETE FROM Evenement_vital WHERE id_habitant=%s", (id,))
        cur.execute("DELETE FROM Mouvement_residuel WHERE id_habitant=%s", (id,))
        cur.execute("DELETE FROM Habitant WHERE id_habitant=%s", (id,))
        conn.commit()
        # Ajouter l'action de suppression a l'historique
        log_action("Suppression habitant", f"ID {id}")

        flash("Habitant supprimé.", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Erreur lors de la suppression : {e}", "danger")
    finally:
        cur.close(); conn.close()
    return redirect(url_for("habitants_list"))

# ─────────────────────────────────────────────────────────────────
# Ménages
# ─────────────────────────────────────────────────────────────────
@app.route("/menages")
@login_required
def menages_list():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    q = request.args.get("q", "").strip()
    secteur = request.args.get("secteur", "").strip()
    revenu_min = request.args.get("revenu_min", "").strip()
    revenu_max = request.args.get("revenu_max", "").strip()
    score_min = request.args.get("score_min", "").strip()
    score_max = request.args.get("score_max", "").strip()
    taille_min = request.args.get("taille_min", "").strip()
    taille_max = request.args.get("taille_max", "").strip()
    sans_chef = request.args.get("sans_chef", "")
    sql = """
        SELECT 
            m.*,
            COUNT(h.id_habitant) AS nb_membres,
            chef.nom AS chef_nom,
            chef.prenom AS chef_prenom
        FROM Menages m
        LEFT JOIN Habitant h ON h.id_menage = m.id_menage
        LEFT JOIN Habitant chef ON chef.id_habitant = m.id_chef_menage
        WHERE 1=1
    """
    params = []
    if q:
        sql += " AND (m.adresse ILIKE %s OR m.secteur ILIKE %s)"
        params += [f"%{q}%", f"%{q}%"]
    if secteur:
        sql += " AND m.secteur = %s"; params.append(secteur)
    if safe_int(revenu_min) is not None:
        sql += " AND m.revenu >= %s"; params.append(safe_int(revenu_min))
    if safe_int(revenu_max) is not None:
        sql += " AND m.revenu <= %s"; params.append(safe_int(revenu_max))
    if safe_int(score_min) is not None:
        sql += " AND COALESCE(m.score_vulnerabilite,0) >= %s"; params.append(safe_int(score_min))
    if safe_int(score_max) is not None:
        sql += " AND COALESCE(m.score_vulnerabilite,0) <= %s"; params.append(safe_int(score_max))
    if sans_chef:
        sql += " AND m.id_chef_menage IS NULL"
    sql += " GROUP BY m.id_menage, chef.nom, chef.prenom"
    having = []
    if safe_int(taille_min) is not None:
        having.append("COUNT(h.id_habitant) >= %s"); params.append(safe_int(taille_min))
    if safe_int(taille_max) is not None:
        having.append("COUNT(h.id_habitant) <= %s"); params.append(safe_int(taille_max))
    if having:
        sql += " HAVING " + " AND ".join(having)
    sql += " ORDER BY m.adresse"
    cur.execute(sql, params)
    menages = cur.fetchall()
    cur.execute("SELECT DISTINCT secteur FROM Menages WHERE COALESCE(secteur,'') <> '' ORDER BY secteur")
    secteurs = [r["secteur"] for r in cur.fetchall()]
    cur.close(); conn.close()
    return render_template("menages_list.html", menages=menages, q=q, secteur=secteur,
                           secteurs=secteurs, revenu_min=revenu_min, revenu_max=revenu_max,
                           score_min=score_min, score_max=score_max, taille_min=taille_min,
                           taille_max=taille_max, sans_chef=sans_chef)

@app.route("/menages/add", methods=["GET", "POST"])
@roles_required("admin", "agent")

def menage_add():
    if request.method == "POST":
        adresse = request.form.get("adresse", "").strip()
        revenu = request.form.get("revenu") or None

        # Nouveaux champs pour la cartographie
        secteur = request.form.get("secteur") or None
        latitude = request.form.get("latitude") or None
        longitude = request.form.get("longitude") or None
        score_vulnerabilite = request.form.get("score_vulnerabilite") or 0
        erreurs = valider_menage(
            adresse,
            revenu,
            latitude,
            longitude,
            score_vulnerabilite
        )

        if erreurs:
            for erreur in erreurs:
                flash(erreur, "danger")
            return render_template("menages_add.html")
        conn = get_db()
        cur = conn.cursor(cursor_factory=RealDictCursor)

        cur.execute("""
            INSERT INTO Menages 
            (adresse, revenu, secteur, latitude, longitude, score_vulnerabilite)
            VALUES (%s, %s, %s, %s, %s, %s)
            RETURNING id_menage
        """, (
            adresse,
            revenu,
            secteur,
            latitude,
            longitude,
            score_vulnerabilite
        ))

        new_id = cur.fetchone()["id_menage"]

        conn.commit()

        # Ajouter l'action d'ajout de ménage à l'historique
        log_action("Ajout ménage", f"Adresse : {adresse} (ID {new_id})")

        flash("Ménage ajouté avec succès.", "success")

        cur.close()
        conn.close()

        return redirect(url_for("menage_detail", id=new_id))

    return render_template("menages_add.html")

@app.route("/carte")
@login_required
def carte():
    return render_template("carte.html")

@app.route("/api/v1/carte/markers", methods=["GET"])
@login_required
def api_carte_markers():
    secteur = request.args.get("secteur", "").strip()

    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    where = [
        "m.latitude IS NOT NULL",
        "m.longitude IS NOT NULL"
    ]
    params = []

    if secteur:
        where.append("m.secteur = %s")
        params.append(secteur)

    where_sql = "WHERE " + " AND ".join(where)

    cur.execute(f"""
        SELECT 
            m.id_menage,
            m.adresse,
            m.latitude,
            m.longitude,
            COALESCE(m.score_vulnerabilite, 0) AS score_vulnerabilite,
            m.secteur,
            COUNT(h.id_habitant) AS nombre_membres
        FROM Menages m
        LEFT JOIN Habitant h ON h.id_menage = m.id_menage
        {where_sql}
        GROUP BY 
            m.id_menage,
            m.adresse,
            m.latitude,
            m.longitude,
            m.score_vulnerabilite,
            m.secteur
        ORDER BY score_vulnerabilite DESC
    """, params)

    menages = cur.fetchall()

    cur.close()
    conn.close()

    markers = []

    for m in menages:
        score = float(m["score_vulnerabilite"] or 0)

        if score >= 70:
            color = "#ef4444"
        elif score >= 50:
            color = "#f59e0b"
        else:
            color = "#10b981"

        markers.append({
            "id": m["id_menage"],
            "name": f"Ménage #{m['id_menage']}",
            "adresse": m["adresse"],
            "lat": float(m["latitude"]),
            "lon": float(m["longitude"]),
            "score_vulnerabilite": score,
            "color": color,
            "secteur": m["secteur"],
            "nombre_membres": int(m["nombre_membres"] or 0)
        })

    return jsonify(markers)


@app.route("/api/v1/secteurs", methods=["GET"])
@login_required
def api_secteurs():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    cur.execute("""
        SELECT DISTINCT secteur
        FROM Menages
        WHERE secteur IS NOT NULL
          AND secteur <> ''
        ORDER BY secteur
    """)

    secteurs = [row["secteur"] for row in cur.fetchall()]

    cur.close()
    conn.close()

    return jsonify({"secteurs": secteurs})

@app.route("/menages/<int:id>")
@login_required
def menage_detail(id):
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("""
        SELECT 
            m.*,
            chef.nom AS chef_nom,
            chef.prenom AS chef_prenom
        FROM Menages m
        LEFT JOIN Habitant chef ON chef.id_habitant = m.id_chef_menage
        WHERE m.id_menage = %s
    """, (id,))
    menage = cur.fetchone()
    if not menage:
        flash("Ménage introuvable.", "danger")
        return redirect(url_for("menages_list"))
    cur.execute("""
        SELECT * FROM Habitant WHERE id_menage=%s ORDER BY nom, prenom
    """, (id,))
    membres = cur.fetchall()
    cur.close(); conn.close()
    return render_template("menages_detail.html", menage=menage, membres=membres)


# Choix du chef de menage 

@app.route("/menages/<int:id>/chef", methods=["POST"])
@roles_required("admin", "agent")
def choisir_chef_menage(id):
    id_chef = request.form.get("id_chef_menage") or None

    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    # Vérifier que le ménage existe
    cur.execute("SELECT id_menage FROM Menages WHERE id_menage = %s", (id,))
    menage = cur.fetchone()

    if not menage:
        flash("Ménage introuvable.", "danger")
        cur.close()
        conn.close()
        return redirect(url_for("menages_list"))

    # Si un chef est choisi, vérifier qu'il appartient bien à ce ménage
    if id_chef:
        cur.execute("""
            SELECT id_habitant, nom, prenom
            FROM Habitant
            WHERE id_habitant = %s
              AND id_menage = %s
              AND statut = 'actif'
        """, (id_chef, id))

        chef = cur.fetchone()

        if not chef:
            flash("Le chef choisi doit être un habitant actif de ce ménage.", "danger")
            cur.close()
            conn.close()
            return redirect(url_for("menage_detail", id=id))

    cur.execute("""
        UPDATE Menages
        SET id_chef_menage = %s
        WHERE id_menage = %s
    """, (id_chef, id))

    if id_chef:
        cur.execute("""
            UPDATE Habitant
            SET lien_avec_chef = NULL
            WHERE id_menage = %s
            AND lien_avec_chef = 'Chef de ménage'
        """, (id,))

        cur.execute("""
            UPDATE Habitant
            SET lien_avec_chef = 'Chef de ménage'
            WHERE id_habitant = %s
            AND id_menage = %s
        """, (id_chef, id))
    conn.commit()

    log_action("Choix chef de ménage", f"Ménage ID {id} — Chef ID {id_chef}")

    flash("Chef de ménage enregistré avec succès.", "success")

    cur.close()
    conn.close()

    return redirect(url_for("menage_detail", id=id))

@app.route("/menages/<int:id>/edit", methods=["GET", "POST"])
@roles_required("admin", "agent")
def menage_edit(id):
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("SELECT * FROM Menages WHERE id_menage=%s", (id,))
    menage = cur.fetchone()
    if not menage:
        flash("Ménage introuvable.", "danger")
        cur.close(); conn.close()
        return redirect(url_for("menages_list"))

    if request.method == "POST":
        adresse             = request.form.get("adresse", "").strip()
        revenu              = request.form.get("revenu") or None
        secteur             = request.form.get("secteur") or None
        latitude            = request.form.get("latitude") or None
        longitude           = request.form.get("longitude") or None
        score_vulnerabilite = request.form.get("score_vulnerabilite") or None

        erreurs = valider_menage(adresse, revenu, latitude, longitude, score_vulnerabilite)
        if erreurs:
            for erreur in erreurs:
                flash(erreur, "danger")
            cur.close(); conn.close()
            return render_template("menages_edit.html", menage=menage)

        cur.execute("""
            UPDATE Menages
            SET adresse=%s, revenu=%s, secteur=%s,
                latitude=%s, longitude=%s, score_vulnerabilite=%s
            WHERE id_menage=%s
        """, (adresse, revenu, secteur, latitude, longitude, score_vulnerabilite, id))
        conn.commit()
        log_action("Modification ménage", f"ID {id} — {adresse}")
        flash("Ménage mis à jour.", "success")
        cur.close(); conn.close()
        return redirect(url_for("menage_detail", id=id))

    cur.close(); conn.close()
    return render_template("menages_edit.html", menage=menage)

@app.route("/menages/<int:id>/delete", methods=["POST"])
@roles_required("admin")
def menage_delete(id):
    conn = get_db()
    cur = conn.cursor()
    try:
        cur.execute("UPDATE Habitant SET id_menage=NULL WHERE id_menage=%s", (id,))
        cur.execute("DELETE FROM Menages WHERE id_menage=%s", (id,))
        conn.commit()
        flash("Ménage supprimé.", "success")
    except Exception as e:
        conn.rollback(); flash(f"Erreur : {e}", "danger")
    finally:
        cur.close(); conn.close()
    return redirect(url_for("menages_list"))







# ─────────────────────────────────────────────────────────────────
# Événements vitaux
# ─────────────────────────────────────────────────────────────────
@app.route("/evenements")
@login_required
def evenements_list():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    type_f = request.args.get("type", "")
    q = request.args.get("q", "").strip()
    date_debut = request.args.get("date_debut", "").strip()
    date_fin = request.args.get("date_fin", "").strip()
    sql = """
        SELECT e.*, h.nom, h.prenom FROM Evenement_vital e
        JOIN Habitant h ON e.id_habitant = h.id_habitant
        WHERE 1=1
    """
    params = []
    if type_f:
        sql += " AND e.type_evenement=%s"; params.append(type_f)
    if q:
        sql += " AND (h.nom ILIKE %s OR h.prenom ILIKE %s OR e.description ILIKE %s)"
        params += [f"%{q}%"] * 3
    if date_debut:
        sql += " AND e.date_evenement >= %s"; params.append(date_debut)
    if date_fin:
        sql += " AND e.date_evenement <= %s"; params.append(date_fin)
    sql += " ORDER BY e.date_evenement DESC"
    cur.execute(sql, params)
    evenements = cur.fetchall()
    cur.close(); conn.close()
    return render_template("evenements_list.html", evenements=evenements, type_f=type_f,
                           q=q, date_debut=date_debut, date_fin=date_fin)

@app.route("/evenements/add", methods=["GET", "POST"])
@roles_required("admin", "agent")
def evenement_add():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    cur.execute("""
        SELECT id_habitant, nom, prenom, statut, date_de_naissance
        FROM Habitant
        WHERE statut <> 'decede'
        ORDER BY nom, prenom
    """)
    habitants = cur.fetchall()

    if request.method == "POST":
        type_evt = request.form.get("type_evenement", "")
        
        # Cas NAISSANCE : créer un nouvel habitant
        if type_evt == "naissance":
            nouveau_prenom = request.form.get("nouveau_prenom", "").strip()
            nouveau_nom = request.form.get("nouveau_nom", "").strip()
            nouveau_genre = request.form.get("nouveau_genre", "").strip()
            nouveau_date_naissance = request.form.get("date_evenement")
            
            # Validation
            if not nouveau_prenom or not nouveau_nom or not nouveau_genre:
                flash("Veuillez remplir tous les champs du nouveau-né (prénom, nom, sexe).", "danger")
                cur.close(); conn.close()
                return render_template("evenements_add.html", habitants=habitants,
                                    today=date.today().strftime('%Y-%m-%d'))
            
            # Normaliser les valeurs
            nouveau_prenom = clean_text(nouveau_prenom)
            nouveau_nom = clean_text(nouveau_nom)
            # Le formulaire envoie directement "Masculin" ou "Féminin"
            
            # Valider le sexe
            if nouveau_genre not in ("Masculin", "Féminin"):
                flash("Sexe invalide. Veuillez choisir Masculin ou Féminin.", "danger")
                cur.close(); conn.close()
                return render_template("evenements_add.html", habitants=habitants,
                                    today=date.today().strftime('%Y-%m-%d'))
            
            # Créer le nouvel habitant
            nouveau_statut = normalize_statut("actif")
            cur.execute("""
                INSERT INTO Habitant (nom, prenom, sexe, date_de_naissance, statut)
                VALUES (%s, %s, %s, %s, %s)
                RETURNING id_habitant
            """, (nouveau_nom, nouveau_prenom, nouveau_genre, nouveau_date_naissance, nouveau_statut))
            
            # Récupérer l'ID du nouvel habitant créé
            nouvel_habitant = cur.fetchone()
            
            if not nouvel_habitant:
                flash("Erreur lors de la création de l'habitant.", "danger")
                cur.close(); conn.close()
                return render_template("evenements_add.html", habitants=habitants,
                                    today=date.today().strftime('%Y-%m-%d'))
            
            d = {
                "type_evenement": type_evt,
                "date_evenement": nouveau_date_naissance,
                "description": request.form.get("description", "").strip(),
                "id_habitant": nouvel_habitant["id_habitant"],
            }
            
            # Enregistrer l'événement de naissance
            cur.execute("""
                INSERT INTO Evenement_vital (type_evenement, date_evenement, description, id_habitant)
                VALUES (%(type_evenement)s, %(date_evenement)s, %(description)s, %(id_habitant)s)
            """, d)
            
            conn.commit()
            cur.close(); conn.close()
            flash("Nouvel habitant enregistré et naissance enregistrée avec succès !", "success")
            return redirect(url_for('evenements_list'))
        
        # Cas DÉCÈS : sélectionner un habitant existant
        d = {
            "type_evenement": type_evt,
            "date_evenement": request.form.get("date_evenement"),
            "description":    request.form.get("description", "").strip(),
            "id_habitant":    request.form.get("id_habitant"),
        }

        if type_evt not in ("naissance", "deces"):
            flash("Type d'événement invalide.", "danger")
            cur.close(); conn.close()
            return render_template("evenements_add.html", habitants=habitants,
                                today=date.today().strftime('%Y-%m-%d'))

        if not d["date_evenement"]:
            flash("Veuillez renseigner la date de l'événement.", "danger")
            cur.close(); conn.close()
            return render_template("evenements_add.html", habitants=habitants,
                                today=date.today().strftime('%Y-%m-%d'))

        if type_evt == "deces" and not d["id_habitant"]:
            flash("Veuillez sélectionner l'habitant concerné par le décès.", "danger")
            cur.close(); conn.close()
            return render_template("evenements_add.html", habitants=habitants,
                                today=date.today().strftime('%Y-%m-%d'))

        cur.execute("""
            SELECT type_mouvement
            FROM Mouvement_residuel
            WHERE id_habitant = %s
        """, (d["id_habitant"],))
        mvt = cur.fetchone()


        # Vérifier l'état de l'habitant avant d'enregistrer l'événement
        cur.execute("""
            SELECT statut, date_de_naissance
            FROM Habitant
            WHERE id_habitant = %s
        """, (d["id_habitant"],))

        h = cur.fetchone()

        if not h:
            flash("Habitant introuvable.", "danger")
            cur.close(); conn.close()
            return render_template("evenements_add.html", habitants=habitants,
                                today=date.today().strftime('%Y-%m-%d'))

        # Vérification DÉCÈS uniquement
        if type_evt == "deces":
            if h["statut"] == "decede":
                flash("Cet habitant est déjà enregistré comme décédé.", "danger")
                cur.close(); conn.close()
                return render_template("evenements_add.html", habitants=habitants,
                                    today=date.today().strftime('%Y-%m-%d'))

            if h["date_de_naissance"] and d["date_evenement"] < str(h["date_de_naissance"]):
                flash("Impossible : la date du décès ne peut pas être antérieure à la date de naissance.", "danger")
                cur.close(); conn.close()
                return render_template("evenements_add.html", habitants=habitants,
                                    today=date.today().strftime('%Y-%m-%d'))
        cur.execute("""
            INSERT INTO Evenement_vital (type_evenement, date_evenement, description, id_habitant)
            VALUES (%(type_evenement)s, %(date_evenement)s, %(description)s, %(id_habitant)s)
        """, d)
        # Mise à jour automatique du statut
        if type_evt == "deces":
            cur.execute("UPDATE Habitant SET statut='decede' WHERE id_habitant=%s",
                        (d["id_habitant"],))
            cur.execute("""
                UPDATE Menages
                SET id_chef_menage = NULL
                WHERE id_chef_menage = %s
            """, (d["id_habitant"],))

            cur.execute("""
                UPDATE Habitant
                SET lien_avec_chef = NULL
                WHERE id_habitant = %s
                AND lien_avec_chef = 'Chef de ménage'
            """, (d["id_habitant"],))
        conn.commit()
        flash("Événement enregistré avec succès.", "success")
        cur.close(); conn.close()
        return redirect(url_for("evenements_list"))

    cur.close(); conn.close()
    return render_template("evenements_add.html", habitants=habitants,
                           today=date.today().strftime('%Y-%m-%d'))

# ─────────────────────────────────────────────────────────────────
# Mouvements résidentiels
# ─────────────────────────────────────────────────────────────────
@app.route("/mouvements")
@login_required
def mouvements_list():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    type_f = request.args.get("type", "")
    q = request.args.get("q", "").strip()
    date_debut = request.args.get("date_debut", "").strip()
    date_fin = request.args.get("date_fin", "").strip()
    sql = """
        SELECT mo.*, h.nom, h.prenom FROM Mouvement_residuel mo
        JOIN Habitant h ON mo.id_habitant = h.id_habitant
        WHERE 1=1
    """
    params = []
    if type_f:
        sql += " AND mo.type_mouvement=%s"; params.append(type_f)
    if q:
        sql += " AND (h.nom ILIKE %s OR h.prenom ILIKE %s OR mo.provenance ILIKE %s OR mo.destination ILIKE %s)"
        params += [f"%{q}%"] * 4
    if date_debut:
        sql += " AND mo.date_mouvement >= %s"; params.append(date_debut)
    if date_fin:
        sql += " AND mo.date_mouvement <= %s"; params.append(date_fin)
    sql += " ORDER BY mo.date_mouvement DESC"
    cur.execute(sql, params)
    mouvements = cur.fetchall()
    cur.close(); conn.close()
    return render_template("mouvements_list.html", mouvements=mouvements, type_f=type_f,
                           q=q, date_debut=date_debut, date_fin=date_fin)

@app.route("/mouvements/add", methods=["GET", "POST"])
@roles_required("admin", "agent")
def mouvement_add():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    cur.execute("SELECT id_habitant, nom, prenom FROM Habitant ORDER BY nom, prenom")
    habitants = cur.fetchall()

    if request.method == "POST":
        type_mvt = request.form.get("type_mouvement", "").strip()
        prov_mvt = request.form.get("provenance", "").strip()
        des_mvt = request.form.get("destination", "").strip()
        date_mouvement = request.form.get("date_mouvement")
        id_habitant = request.form.get("id_habitant")

        # Vérification des champs obligatoires généraux
        if not type_mvt:
            flash("Veuillez sélectionner le type de mouvement.", "danger")
            cur.close()
            conn.close()
            return render_template("mouvements_add.html", habitants=habitants)

        if not date_mouvement:
            flash("Veuillez renseigner la date du mouvement.", "danger")
            cur.close()
            conn.close()
            return render_template("mouvements_add.html", habitants=habitants)

        if not id_habitant:
            flash("Veuillez sélectionner un habitant.", "danger")
            cur.close()
            conn.close()
            return render_template("mouvements_add.html", habitants=habitants)

        # Vérification selon le type de mouvement
        if type_mvt == "arrivee" and not prov_mvt:
            flash("Veuillez renseigner le lieu de provenance pour une arrivée.", "danger")
            cur.close()
            conn.close()
            return render_template("mouvements_add.html", habitants=habitants)

        if type_mvt == "depart" and not des_mvt:
            flash("Veuillez renseigner la destination pour un départ.", "danger")
            cur.close()
            conn.close()
            return render_template("mouvements_add.html", habitants=habitants)

        if type_mvt == "interne" and not des_mvt:
            flash("Veuillez renseigner la destination ou la nouvelle localisation pour un mouvement interne.", "danger")
            cur.close()
            conn.close()
            return render_template("mouvements_add.html", habitants=habitants)

        # Transformer les chaînes vides en None avant insertion
        prov_mvt = prov_mvt if prov_mvt else None
        des_mvt = des_mvt if des_mvt else None

        d = {
            "type_mouvement": type_mvt,
            "date_mouvement": date_mouvement,
            "provenance": prov_mvt,
            "destination": des_mvt,
            "id_habitant": id_habitant,
        }

        cur.execute("SELECT statut FROM Habitant WHERE id_habitant=%s", (id_habitant,))
        h = cur.fetchone()

        if h:
            if type_mvt == "depart" and h["statut"] in ("decede", "parti"):
                flash("Cet habitant est déjà décédé ou parti.", "danger")
                cur.close()
                conn.close()
                return render_template("mouvements_add.html", habitants=habitants)

            if type_mvt == "arrivee" and h["statut"] == "decede":
                flash("Cet habitant est décédé, une arrivée est impossible.", "danger")
                cur.close()
                conn.close()
                return render_template("mouvements_add.html", habitants=habitants)

        cur.execute("""
            INSERT INTO Mouvement_residuel
              (type_mouvement, date_mouvement, provenance, destination, id_habitant)
            VALUES
              (%(type_mouvement)s, %(date_mouvement)s, %(provenance)s, %(destination)s, %(id_habitant)s)
        """, d)

        if type_mvt == "depart":
            cur.execute("""
                UPDATE Habitant 
                SET statut='parti' 
                WHERE id_habitant=%s
            """, (id_habitant,))

            cur.execute("""
                UPDATE Menages
                SET id_chef_menage = NULL
                WHERE id_chef_menage = %s
            """, (id_habitant,))

            cur.execute("""
                UPDATE Habitant
                SET lien_avec_chef = NULL
                WHERE id_habitant = %s
                AND lien_avec_chef = 'Chef de ménage'
            """, (id_habitant,))

        elif type_mvt == "arrivee":
            cur.execute("""
                UPDATE Habitant 
                SET statut='actif' 
                WHERE id_habitant=%s
            """, (id_habitant,))

        conn.commit()

        flash("Mouvement enregistré avec succès.", "success")

        cur.close()
        conn.close()

        return redirect(url_for("mouvements_list"))

    cur.close()
    conn.close()

    return render_template(
        "mouvements_add.html",
        habitants=habitants,
        today=date.today().strftime('%Y-%m-%d')
    )


@app.route("/import/habitants", methods=["GET", "POST"])
@roles_required("admin", "agent")
def import_habitants():
    result = None
    expected_columns = [
        "nom", "prenom", "sexe", "date_de_naissance", "profession",
        "niveau_etude", "telephone", "statut", "id_menage", "adresse_menage",
        "lien_avec_chef"
    ]

    if request.method == "POST":
        upload = request.files.get("fichier")
        if not upload or not upload.filename:
            flash("Veuillez sélectionner un fichier CSV ou Excel.", "danger")
            return render_template("import_habitants.html", expected_columns=expected_columns, result=result)

        filename = secure_filename(upload.filename)
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in ("csv", "xlsx", "xls"):
            flash("Format non supporté. Utilisez CSV, XLSX ou XLS.", "danger")
            return render_template("import_habitants.html", expected_columns=expected_columns, result=result)

        try:
            if ext == "csv":
                df = pd.read_csv(upload, dtype=str).fillna("")
            else:
                df = pd.read_excel(upload, dtype=str).fillna("")
        except Exception as e:
            flash(f"Lecture du fichier impossible : {e}", "danger")
            return render_template("import_habitants.html", expected_columns=expected_columns, result=result)

        if len(df.index) > 1000:
            flash("Import limité à 1000 lignes par opération pour garder un traitement fiable.", "danger")
            return render_template("import_habitants.html", expected_columns=expected_columns, result=result)

        column_map = {normalize_key(col): col for col in df.columns}

        def col(*names):
            for name in names:
                key = normalize_key(name)
                if key in column_map:
                    return column_map[key]
            return None

        conn = get_db()
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute("SELECT id_menage, adresse FROM Menages")
        menages_by_id = {str(row["id_menage"]): row["id_menage"] for row in cur.fetchall()}
        cur.execute("SELECT id_menage, LOWER(TRIM(adresse)) adresse FROM Menages")
        menages_by_adresse = {row["adresse"]: row["id_menage"] for row in cur.fetchall()}

        inserted = 0
        skipped = 0
        errors = []
        for idx, row in df.iterrows():
            line_no = idx + 2
            nom = clean_text(row.get(col("nom", "name"), "")).upper()
            prenom = clean_text(row.get(col("prenom", "prénom", "first_name"), ""))
            sexe = normalize_sexe(row.get(col("sexe", "genre"), ""))
            statut = normalize_statut(row.get(col("statut", "status"), "actif"))
            date_naissance = parse_import_date(row.get(col("date_de_naissance", "naissance", "date naissance"), ""))
            telephone = clean_text(row.get(col("telephone", "téléphone", "tel"), ""))
            profession = clean_text(row.get(col("profession", "metier", "métier"), ""))
            niveau = clean_text(row.get(col("niveau_etude", "niveau d'etude", "niveau d'étude"), ""))
            lien = clean_text(row.get(col("lien_avec_chef", "lien chef", "position_menage"), ""))
            id_menage_value = clean_text(row.get(col("id_menage", "menage", "ménage"), ""))
            adresse_menage = clean_text(row.get(col("adresse_menage", "adresse"), "")).lower()

            row_errors = []
            if not nom:
                row_errors.append("nom manquant")
            if not prenom:
                row_errors.append("prénom manquant")
            if sexe not in ("Masculin", "Féminin"):
                row_errors.append("sexe invalide")
            if statut not in ("actif", "decede", "parti"):
                row_errors.append("statut invalide")
            if date_naissance and date_naissance > date.today():
                row_errors.append("date de naissance future")
            if telephone and not re.match(r"^[+0-9 ()-]{8,20}$", telephone):
                row_errors.append("téléphone invalide")

            id_menage = None
            if id_menage_value:
                id_menage = menages_by_id.get(id_menage_value)
                if not id_menage:
                    row_errors.append("id_menage introuvable")
            elif adresse_menage:
                id_menage = menages_by_adresse.get(adresse_menage)
                if not id_menage:
                    row_errors.append("adresse_menage introuvable")

            if row_errors:
                skipped += 1
                errors.append(f"Ligne {line_no}: {', '.join(row_errors)}")
                continue

            cur.execute("""
                SELECT id_habitant FROM Habitant
                WHERE UPPER(TRIM(nom))=%s
                  AND UPPER(TRIM(prenom))=%s
                  AND (date_de_naissance IS NOT DISTINCT FROM %s)
                LIMIT 1
            """, (nom, prenom.upper(), date_naissance))
            if cur.fetchone():
                skipped += 1
                errors.append(f"Ligne {line_no}: doublon potentiel ignoré")
                continue

            cur.execute("SAVEPOINT import_row")
            try:
                cur.execute("""
                    INSERT INTO Habitant
                      (nom, prenom, sexe, date_de_naissance, profession, niveau_etude,
                       telephone, statut, id_menage, lien_avec_chef)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                    RETURNING id_habitant
                """, (
                    nom, prenom, sexe, date_naissance, profession, niveau,
                    telephone, statut, id_menage, lien or None
                ))
                new_id = cur.fetchone()["id_habitant"]
                if lien == "Chef de ménage" and id_menage:
                    cur.execute("UPDATE Menages SET id_chef_menage=%s WHERE id_menage=%s", (new_id, id_menage))
                cur.execute("RELEASE SAVEPOINT import_row")
                inserted += 1
            except Exception as e:
                skipped += 1
                cur.execute("ROLLBACK TO SAVEPOINT import_row")
                cur.execute("RELEASE SAVEPOINT import_row")
                errors.append(f"Ligne {line_no}: {e}")

        conn.commit()
        cur.close()
        conn.close()
        log_action("Import habitants", f"{inserted} ajout(s), {skipped} ligne(s) ignorée(s)")
        result = {"inserted": inserted, "skipped": skipped, "errors": errors[:30], "total": len(df.index)}
        flash(f"Import terminé : {inserted} habitant(s) ajouté(s), {skipped} ligne(s) ignorée(s).", "success")

    return render_template("import_habitants.html", expected_columns=expected_columns, result=result)


@app.route("/qualite")
@roles_required("admin", "responsable", "agent")
def qualite_donnees():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    snapshot = build_quality_snapshot(cur)
    cur.close()
    conn.close()
    return render_template("qualite.html", snapshot=snapshot)


def generer_rapport_narratif(stats):
    pop_active = stats.get("pop_active", 0) or 0
    pop_totale = stats.get("pop_totale", 0) or 0
    nb_menages = stats.get("nb_menages", 0) or 0
    naissances = stats.get("total_naissances", 0) or 0
    deces = stats.get("total_deces", 0) or 0
    arrivees = stats.get("total_arrivees", 0) or 0
    departs = stats.get("total_departs", 0) or 0
    nb_decedes = stats.get("nb_decedes", 0) or 0
    nb_partis = stats.get("nb_partis", 0) or 0
    taille_moy = stats.get("taille_moy_menage", 0) or 0

    solde_migratoire = arrivees - departs

    return f"""
Le présent rapport statistique présente une analyse synthétique de la situation démographique du quartier de Danga, à partir des données enregistrées dans l’Annuaire Statistique local.

À la date de génération du rapport, le quartier compte {pop_active} habitants actifs sur une population totale enregistrée de {pop_totale} habitants. Cette population est répartie dans {nb_menages} ménages, avec une taille moyenne estimée à {taille_moy} personne(s) par ménage.

L’analyse des événements vitaux indique que {naissances} naissance(s) et {deces} décès ont été enregistrés dans la base. Le nombre total de personnes déclarées décédées est de {nb_decedes}. Ces indicateurs permettent de suivre l’évolution naturelle de la population et d’apprécier les changements démographiques observés dans le quartier.

Concernant les mouvements résidentiels, la base fait ressortir {arrivees} arrivée(s) contre {departs} départ(s), soit un solde migratoire de {solde_migratoire}. Cet indicateur permet d’apprécier l’attractivité résidentielle du quartier et la mobilité de sa population.

Par ailleurs, {nb_partis} habitant(s) sont actuellement déclarés comme ayant quitté le quartier. Cette information complète l’analyse de la dynamique résidentielle et permet de mieux comprendre l’évolution de la composition de la population locale.

Dans l’ensemble, les résultats présentés constituent une base utile pour le suivi démographique, la planification locale, l’orientation des actions communautaires et l’aide à la décision des responsables locaux.
""".strip()
# ─────────────────────────────────────────────────────────────────
# Rapports & statistiques
# ─────────────────────────────────────────────────────────────────
@app.route("/rapports")
@login_required
def rapports():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    stats = collect_report_stats(cur)

    cur.close()
    conn.close()

    rapport_narratif = generer_rapport_narratif(stats)

    return render_template(
    "rapports.html",
    rapport_narratif=rapport_narratif,
    today=date.today().strftime("%d/%m/%Y")
)


@app.route("/rapports/impression")
@login_required
def rapport_impression():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    stats = collect_report_stats(cur)
    quality = build_quality_snapshot(cur)
    cur.close()
    conn.close()
    return render_template(
        "rapport_impression.html",
        stats=stats,
        quality=quality,
        rapport_narratif=generer_rapport_narratif(stats),
        today=date.today().strftime("%d/%m/%Y")
    )


@app.route("/export/rapport_word")
@login_required
def export_rapport_word():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    stats = {}

    cur.execute("SELECT COUNT(*) c FROM Habitant WHERE statut='actif'")
    stats["pop_active"] = cur.fetchone()["c"]

    cur.execute("SELECT COUNT(*) c FROM Habitant")
    stats["pop_totale"] = cur.fetchone()["c"]

    cur.execute("SELECT COUNT(*) c FROM Habitant WHERE statut='decede'")
    stats["nb_decedes"] = cur.fetchone()["c"]

    cur.execute("SELECT COUNT(*) c FROM Habitant WHERE statut='parti'")
    stats["nb_partis"] = cur.fetchone()["c"]

    cur.execute("SELECT COUNT(*) c FROM Menages")
    stats["nb_menages"] = cur.fetchone()["c"]

    cur.execute("SELECT COUNT(*) c FROM Evenement_vital WHERE type_evenement='naissance'")
    stats["total_naissances"] = cur.fetchone()["c"]

    cur.execute("SELECT COUNT(*) c FROM Evenement_vital WHERE type_evenement='deces'")
    stats["total_deces"] = cur.fetchone()["c"]

    cur.execute("SELECT COUNT(*) c FROM Mouvement_residuel WHERE type_mouvement='arrivee'")
    stats["total_arrivees"] = cur.fetchone()["c"]

    cur.execute("SELECT COUNT(*) c FROM Mouvement_residuel WHERE type_mouvement='depart'")
    stats["total_departs"] = cur.fetchone()["c"]

    cur.execute("""
        SELECT ROUND(AVG(nb),2) as moy FROM
        (SELECT COUNT(*) nb FROM Habitant
        WHERE statut = 'actif' GROUP BY id_menage) t
    """)
    row = cur.fetchone()
    stats["taille_moy_menage"] = row["moy"] if row and row["moy"] else 0

    cur.close()
    conn.close()

    texte = generer_rapport_narratif(stats)

    doc = Document()
    titre = doc.add_heading("Rapport statistique du quartier de Danga", level=1)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for paragraphe in texte.split("\n\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(paragraphe.strip())
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"rapport_narratif_{date.today():%Y%m%d}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
# ─────────────────────────────────────────────────────────────────
# API REST documentee
# ─────────────────────────────────────────────────────────────────
@app.route("/api/docs")
@login_required
def api_docs():
    return render_template("api_docs.html")


@app.route("/api/v1/stats")
@login_required
def api_v1_stats():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    stats = collect_report_stats(cur)
    quality = build_quality_snapshot(cur)
    cur.close()
    conn.close()
    return jsonify({
        "stats": stats,
        "quality": {
            "score": quality["score"],
            "issue_total": quality["issue_total"],
            "metrics": quality["metrics"],
        }
    })


@app.route("/api/v1/habitants")
@login_required
def api_v1_habitants():
    q = request.args.get("q", "").strip()
    limit = min(safe_int(request.args.get("limit")) or 100, 200)
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    sql = """
        SELECT h.id_habitant, h.nom, h.prenom, h.sexe, h.date_de_naissance,
               h.profession, h.niveau_etude, h.telephone, h.statut,
               m.adresse, m.secteur
        FROM Habitant h
        LEFT JOIN Menages m ON m.id_menage = h.id_menage
        WHERE 1=1
    """
    params = []
    if q:
        sql += " AND (h.nom ILIKE %s OR h.prenom ILIKE %s OR h.telephone ILIKE %s OR m.adresse ILIKE %s)"
        params += [f"%{q}%"] * 4
    sql += " ORDER BY h.nom, h.prenom LIMIT %s"; params.append(limit)
    cur.execute(sql, params)
    rows = [dict(row) for row in cur.fetchall()]
    cur.close()
    conn.close()
    return jsonify({"count": len(rows), "items": rows})


@app.route("/api/v1/menages")
@login_required
def api_v1_menages():
    q = request.args.get("q", "").strip()
    limit = min(safe_int(request.args.get("limit")) or 100, 200)
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    sql = """
        SELECT m.id_menage, m.adresse, m.secteur, m.revenu,
               m.latitude, m.longitude, m.score_vulnerabilite,
               COUNT(h.id_habitant) nb_membres
        FROM Menages m
        LEFT JOIN Habitant h ON h.id_menage = m.id_menage
        WHERE 1=1
    """
    params = []
    if q:
        sql += " AND (m.adresse ILIKE %s OR m.secteur ILIKE %s)"
        params += [f"%{q}%", f"%{q}%"]
    sql += " GROUP BY m.id_menage ORDER BY m.adresse LIMIT %s"; params.append(limit)
    cur.execute(sql, params)
    rows = [dict(row) for row in cur.fetchall()]
    cur.close()
    conn.close()
    return jsonify({"count": len(rows), "items": rows})


# ─────────────────────────────────────────────────────────────────
# API JSON (pour les graphiques JS)
# ─────────────────────────────────────────────────────────────────
@app.route("/api/charts")
@login_required
def api_charts():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)

    cur.execute("""
        SELECT COALESCE(NULLIF(sexe,''),'Inconnu') as sexe, COUNT(*) cnt
        FROM Habitant WHERE statut='actif' GROUP BY sexe
    """)
    sexe = [dict(r) for r in cur.fetchall()]

    cur.execute("""
        SELECT
            CASE
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 0  AND 9  THEN '0–9'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 10 AND 19 THEN '10–19'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 20 AND 29 THEN '20–29'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 30 AND 39 THEN '30–39'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 40 AND 49 THEN '40–49'
                WHEN EXTRACT(YEAR FROM AGE(date_de_naissance)) BETWEEN 50 AND 59 THEN '50–59'
                ELSE '60+'
            END as tranche, COUNT(*) cnt
        FROM Habitant WHERE date_de_naissance IS NOT NULL AND statut = 'actif'
        GROUP BY tranche ORDER BY tranche
    """)

    ages = [dict(r) for r in cur.fetchall()]

    cur.execute("""
        SELECT
            TO_CHAR(DATE_TRUNC('month', date_evenement), 'Mon YYYY') as mois,
            DATE_TRUNC('month', date_evenement) as ord,
            SUM(CASE WHEN type_evenement='naissance' THEN 1 ELSE 0 END) naissances,
            SUM(CASE WHEN type_evenement='deces' THEN 1 ELSE 0 END) deces
        FROM Evenement_vital
        WHERE date_evenement >= CURRENT_DATE - INTERVAL '12 months'
        GROUP BY DATE_TRUNC('month', date_evenement)
        ORDER BY ord
    """)
    rows = cur.fetchall()
    evolution = [{"mois": r["mois"], "naissances": r["naissances"], "deces": r["deces"]}
                 for r in rows]

    cur.close(); conn.close()
    return jsonify({"sexe": sexe, "ages": ages, "evolution": evolution})

# ─────────────────────────────────────────────────────────────────
# Historique des actions
# ─────────────────────────────────────────────────────────────────
@app.route("/historique")
@roles_required("admin", "responsable")
def historique():
    conn = get_db()
    cur  = conn.cursor(cursor_factory=RealDictCursor)

    # Filtres
    utilisateur = request.args.get("utilisateur", "")
    date_debut  = request.args.get("date_debut", "")
    date_fin    = request.args.get("date_fin", "")
    action_f    = request.args.get("action", "")

    sql    = "SELECT * FROM historique_actions WHERE 1=1"
    params = []
    if utilisateur:
        sql += " AND utilisateur = %s"; params.append(utilisateur)
    if date_debut:
        sql += " AND date_action >= %s"; params.append(date_debut)
    if date_fin:
        sql += " AND date_action <= %s + INTERVAL '1 day'"; params.append(date_fin)
    if action_f:
        sql += " AND action ILIKE %s"; params.append(f"%{action_f}%")
    sql += " ORDER BY date_action DESC LIMIT 500"

    cur.execute(sql, params)
    logs = cur.fetchall()

    cur.execute("SELECT DISTINCT utilisateur FROM historique_actions ORDER BY utilisateur")
    utilisateurs = [r["utilisateur"] for r in cur.fetchall()]

    cur.close(); conn.close()
    return render_template("historique.html", logs=logs,
                           utilisateurs=utilisateurs,
                           utilisateur=utilisateur, date_debut=date_debut,
                           date_fin=date_fin, action_f=action_f)


# ─────────────────────────────────────────────────────────────────
# Paramètres (admin uniquement)
# ─────────────────────────────────────────────────────────────────
@app.route("/parametres", methods=["GET", "POST"])
@roles_required("admin")
def parametres():
    if request.method == "POST":
        action = request.form.get("action")

        if action == "change_password":
            username = request.form.get("role")
            new_pw   = request.form.get("new_password", "")
            confirm  = request.form.get("confirm_password", "")

            conn = get_db()
            cur = conn.cursor(cursor_factory=RealDictCursor)
            cur.execute("SELECT * FROM utilisateurs WHERE username=%s", (username,))
            user_row = cur.fetchone()

            if not user_row:
                flash("Utilisateur introuvable.", "danger")
            elif new_pw != confirm:
                flash("Les mots de passe ne correspondent pas.", "danger")
            elif validate_password_strength(new_pw):
                for error in validate_password_strength(new_pw):
                    flash(error, "danger")
            else:
                hashed = bcrypt.hashpw(new_pw.encode(), bcrypt.gensalt()).decode()
                cur.execute(
                    "UPDATE utilisateurs SET password=%s WHERE username=%s",
                    (hashed, username)
                )
                conn.commit()
                log_action("Modification mot de passe", f"Profil modifié : {username}")
                flash("Mot de passe mis à jour.", "success")

            cur.close()
            conn.close()

        elif action == "add_user":
            new_username = request.form.get("new_role", "").strip().lower()
            new_label    = request.form.get("new_label", "").strip()
            new_pw       = request.form.get("new_user_password", "")
            new_email    = request.form.get("new_email", "").strip()
            user_role    = request.form.get("user_role", "agent")

            if not new_username or not new_label or not new_pw or not new_email:
                flash("Tous les champs sont obligatoires.", "danger")
            elif validate_password_strength(new_pw):
                for error in validate_password_strength(new_pw):
                    flash(error, "danger")
            else:
                hashed = bcrypt.hashpw(new_pw.encode(), bcrypt.gensalt()).decode()
                conn = get_db()
                cur = conn.cursor()

                try:
                    cur.execute("""
                        INSERT INTO utilisateurs 
                        (username, password, role, label, email)
                        VALUES (%s, %s, %s, %s, %s)
                    """, (new_username, hashed, user_role, new_label, new_email))
                    conn.commit()
                    log_action("Ajout utilisateur", f"Nouveau profil : {new_username} ({user_role})")
                    flash(f"Utilisateur « {new_label} » ajouté.", "success")
                except Exception as e:
                    conn.rollback()
                    flash(f"Erreur : {e}", "danger")
                finally:
                    cur.close()
                    conn.close()

        elif action == "change_email":
            username  = request.form.get("email_role", "").strip()
            new_email = request.form.get("new_email_value", "").strip()

            if not username or not new_email:
                flash("Tous les champs sont obligatoires.", "danger")
            elif "@" not in new_email:
                flash("Adresse email invalide.", "danger")
            else:
                conn = get_db()
                cur = conn.cursor()
                cur.execute("UPDATE utilisateurs SET email=%s WHERE username=%s", (new_email, username))
                conn.commit()
                cur.close(); conn.close()
                log_action("Modification email récupération", f"Profil : {username} → {new_email}")
                flash(f"Email de récupération mis à jour pour « {username} ».", "success")

        elif action == "delete_user":
            username = request.form.get("role")

            if username == "admin":
                flash("Impossible de supprimer l'administrateur.", "danger")
            else:
                conn = get_db()
                cur = conn.cursor()
                cur.execute("DELETE FROM utilisateurs WHERE username=%s", (username,))
                conn.commit()
                cur.close()
                conn.close()

                log_action("Suppression utilisateur", f"Profil supprimé : {username}")
                flash("Utilisateur supprimé.", "success")

        return redirect(url_for("parametres"))

    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("SELECT * FROM utilisateurs ORDER BY role, username")
    users = {u["username"]: u for u in cur.fetchall()}
    cur.close()
    conn.close()

    return render_template("parametres.html", users=users)
# ─────────────────────────────────────────────────────────────────
# Exports Excel
# ─────────────────────────────────────────────────────────────────
@app.route("/export/habitants")
@roles_required("admin", "responsable")
def export_habitants():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute("""
        SELECT h.id_habitant "ID", h.nom "Nom", h.prenom "Prénom", h.sexe "Sexe",
               h.date_de_naissance "Date de naissance", h.profession "Profession",
               h.niveau_etude "Niveau d'étude", h.telephone "Téléphone",
               h.statut "Statut", m.adresse "Adresse ménage"
        FROM Habitant h LEFT JOIN Menages m ON h.id_menage = m.id_menage
        ORDER BY h.nom, h.prenom
    """)
    df = pd.DataFrame(cur.fetchall())
    cur.close(); conn.close()
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Habitants", index=False)
    buf.seek(0)
    return send_file(buf,
                     download_name=f"habitants_{date.today():%Y%m%d}.xlsx",
                     as_attachment=True,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/export/rapport_complet")
@roles_required("admin", "responsable")
def export_rapport():
    conn = get_db()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        # Habitants
        cur.execute("""
            SELECT h.nom "Nom", h.prenom "Prénom", h.sexe "Sexe",
                   h.date_de_naissance "Naissance", h.profession "Profession",
                   h.statut "Statut", m.adresse "Adresse"
            FROM Habitant h LEFT JOIN Menages m ON h.id_menage=m.id_menage ORDER BY h.nom
        """)
        pd.DataFrame(cur.fetchall()).to_excel(w, sheet_name="Habitants", index=False)
        # Ménages
        cur.execute("""
            SELECT m.id_menage "ID", m.adresse "Adresse", m.revenu "Revenu (FCFA)",
                   COUNT(h.id_habitant) "Nb membres"
            FROM Menages m LEFT JOIN Habitant h ON h.id_menage=m.id_menage
            GROUP BY m.id_menage ORDER BY m.adresse
        """)
        pd.DataFrame(cur.fetchall()).to_excel(w, sheet_name="Ménages", index=False)
        # Événements
        cur.execute("""
            SELECT e.type_evenement "Type", e.date_evenement "Date",
                   h.nom "Nom", h.prenom "Prénom", e.description "Description"
            FROM Evenement_vital e JOIN Habitant h ON e.id_habitant=h.id_habitant
            ORDER BY e.date_evenement DESC
        """)
        pd.DataFrame(cur.fetchall()).to_excel(w, sheet_name="Événements", index=False)
        # Mouvements
        cur.execute("""
            SELECT mo.type_mouvement "Type", mo.date_mouvement "Date",
                   h.nom "Nom", h.prenom "Prénom",
                   mo.provenance "Provenance", mo.destination "Destination"
            FROM Mouvement_residuel mo JOIN Habitant h ON mo.id_habitant=h.id_habitant
            ORDER BY mo.date_mouvement DESC
        """)
        pd.DataFrame(cur.fetchall()).to_excel(w, sheet_name="Mouvements", index=False)
        # Statistiques résumé
        stats_items = []
        for label, sql in [
            ("Population active",    "SELECT COUNT(*) c FROM Habitant WHERE statut='actif'"),
            ("Population totale",    "SELECT COUNT(*) c FROM Habitant"),
            ("Nombre de ménages",    "SELECT COUNT(*) c FROM Menages"),
            ("Total naissances",     "SELECT COUNT(*) c FROM Evenement_vital WHERE type_evenement='naissance'"),
            ("Total décès",          "SELECT COUNT(*) c FROM Evenement_vital WHERE type_evenement='deces'"),
            ("Total arrivées",       "SELECT COUNT(*) c FROM Mouvement_residuel WHERE type_mouvement='arrivee'"),
            ("Total départs",        "SELECT COUNT(*) c FROM Mouvement_residuel WHERE type_mouvement='depart'"),
        ]:
            cur.execute(sql)
            stats_items.append({"Indicateur": label, "Valeur": cur.fetchone()["c"]})
        pd.DataFrame(stats_items).to_excel(w, sheet_name="Statistiques", index=False)
    cur.close(); conn.close()
    buf.seek(0)
    return send_file(buf,
                     download_name=f"rapport_quartier_{date.today():%Y%m%d}.xlsx",
                     as_attachment=True,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.errorhandler(404)
def page_not_found(e):
    return render_template("404.html"), 404

@app.errorhandler(500)
def internal_error(e):
    return render_template("500.html"), 500

if __name__ == "__main__":
    app.run(debug=True, use_reloader=False)
